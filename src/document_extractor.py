from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from pypdf import PdfReader
from pptx import Presentation


SUPPORTED_EXTENSIONS = {".pdf", ".xlsx", ".xls", ".pptx"}


def list_documents(data_dir: Path) -> list[Path]:
    """Return supported diligence-room files from the Data directory."""
    if not data_dir.exists():
        return []

    files = [
        path
        for path in data_dir.iterdir()
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_EXTENSIONS
        and not path.name.startswith("~$")
    ]

    return sorted(files, key=lambda p: p.name.lower())


def extract_pdf(path: Path) -> dict[str, Any]:
    """Extract text from a PDF page by page."""
    reader = PdfReader(str(path))

    pages: list[dict[str, Any]] = []
    full_text_parts: list[str] = []

    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            text = f"[PDF extraction error: {exc}]"

        text = text.strip()

        pages.append(
            {
                "page": page_number,
                "text": text,
            }
        )

        if text:
            full_text_parts.append(
                f"[PAGE {page_number}]\n{text}"
            )

    return {
        "file_name": path.name,
        "file_type": "pdf",
        "text": "\n\n".join(full_text_parts),
        "pages": pages,
    }


def extract_pptx(path: Path) -> dict[str, Any]:
    """Extract text from a PowerPoint presentation slide by slide."""
    presentation = Presentation(str(path))

    slides: list[dict[str, Any]] = []
    full_text_parts: list[str] = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        slide_text_parts: list[str] = []

        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue

            try:
                text = shape.text.strip()
            except Exception:
                text = ""

            if text:
                slide_text_parts.append(text)

        slide_text = "\n".join(slide_text_parts)

        slides.append(
            {
                "slide": slide_number,
                "text": slide_text,
            }
        )

        if slide_text:
            full_text_parts.append(
                f"[SLIDE {slide_number}]\n{slide_text}"
            )

    return {
        "file_name": path.name,
        "file_type": "pptx",
        "text": "\n\n".join(full_text_parts),
        "slides": slides,
    }


def extract_excel(path: Path) -> dict[str, Any]:
    """Extract worksheets from an Excel workbook."""
    workbook = pd.ExcelFile(path)

    sheets: dict[str, Any] = {}
    full_text_parts: list[str] = []

    for sheet_name in workbook.sheet_names:
        dataframe = pd.read_excel(
            path,
            sheet_name=sheet_name,
            header=None,
        ).fillna("")

        sheets[sheet_name] = dataframe

        rows: list[str] = []

        for row in dataframe.itertuples(index=False, name=None):
            values = [str(value).strip() for value in row]

            if not any(values):
                continue

            rows.append(" | ".join(values))

        if rows:
            full_text_parts.append(
                f"[SHEET {sheet_name}]\n"
                + "\n".join(rows)
            )

    return {
        "file_name": path.name,
        "file_type": "xlsx",
        "text": "\n\n".join(full_text_parts),
        "sheets": sheets,
    }


def extract_document(path: Path) -> dict[str, Any]:
    """Extract a supported document based on its extension."""
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return extract_pdf(path)

    if suffix == ".pptx":
        return extract_pptx(path)

    if suffix in {".xlsx", ".xls"}:
        return extract_excel(path)

    raise ValueError(
        f"Unsupported file type: {path.suffix}"
    )


def extract_all_documents(
    data_dir: Path,
) -> list[dict[str, Any]]:
    """Extract all supported documents from the data room."""
    documents: list[dict[str, Any]] = []

    for path in list_documents(data_dir):
        try:
            documents.append(
                extract_document(path)
            )
        except Exception as exc:
            documents.append(
                {
                    "file_name": path.name,
                    "file_type": path.suffix.lower().lstrip("."),
                    "text": "",
                    "error": str(exc),
                }
            )

    return documents